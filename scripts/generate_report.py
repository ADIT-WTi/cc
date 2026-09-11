# ============================================================
#         CRD BUSINESS REPORT GENERATOR - GITHUB ACTIONS VERSION
# ============================================================
# What this script does:
#   1. Reads state file from process_invoice.py
#   2. Loads combined invoice data
#   3. Filters CRD vertical (excludes Not Found rows)
#   4. Generates professional FY-aware Excel business report
#   5. Saves report to reports/YYYY/MonthName/
#   6. Updates state file with report path
# ============================================================

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import numpy as np
import os
import json
from datetime import datetime

# ============================================================
#                    PATH CONFIGURATION
# ============================================================

def get_repo_root():
    """Returns the root directory of the repository."""
    return os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def read_state_file():
    """Reads state file and resolves all paths."""
    repo_root  = get_repo_root()
    state_path = os.path.join(
        repo_root, "reports", ".run_state.json"
    )

    if not os.path.exists(state_path):
        raise FileNotFoundError(
            f"❌ State file not found: {state_path}"
        )

    with open(state_path, "r") as f:
        state = json.load(f)

    path_keys = [
        "output_dir",
        "not_found_summary_path",
        "not_found_json_path",
        "combined_file_path",
        "sample_invoice_path",
        "final_report_path",
        "final_sample_invoice_path"
    ]

    for key in path_keys:
        if key in state and state[key]:
            state[key] = resolve_path(state[key])

    return state, state_path


def update_state_file(state_path, updates):
    with open(state_path, "r") as f:
        state = json.load(f)
    state.update(updates)
    with open(state_path, "w") as f:
        json.dump(state, f, indent=4)
    return state


# ============================================================
#                    STYLE HELPERS
# ============================================================

DARK_NAVY    = "0D1B2A"
NAVY_3       = "1B3A5C"
HEADER_BG    = "2E5090"
LIGHT_HEADER = "3A6BC5"
SUB_BOOK_BG  = "4472C4"
SUB_REV_BG   = "2E75B6"
SUMMARY_BG   = "1F4E79"
OVERALL_BG   = "8B4513"
GT_COL_BG    = "FFF2CC"
ACCENT_GOLD  = "FFD700"
ACCENT_AMBER = "FFC000"
ALT_ROW_1    = "F7FAFF"
ALT_ROW_2    = "E8F0FE"
WHITE        = "FFFFFF"
BLACK        = "1A1A1A"
GREY_DARK    = "404040"
GREY_MED     = "808080"
BORDER_BLUE  = "B4C6E7"
BORDER_DARK  = "4472C4"
DATA_BOOK_BG1 = "F7FAFF"
DATA_BOOK_BG2 = "EBF1FB"
DATA_REV_BG1  = "EFF6FF"
DATA_REV_BG2  = "E0ECFA"
FY_SUMMARY_BG1   = "FFF9E0"
FY_SUMMARY_BG2   = "FFF4C4"
OVERALL_COL_BG1  = "F5E6D3"
OVERALL_COL_BG2  = "EDD9BE"
GT_FONT_DARK  = "1B2A4A"
SUBTITLE_BG   = "2E5090"
FY_SEP_COLOR  = "1F4E79"


def make_fill(c):
    return PatternFill(start_color=c, end_color=c, fill_type='solid')


def make_border(c, s='thin'):
    return Border(
        left   = Side(style=s, color=c),
        right  = Side(style=s, color=c),
        top    = Side(style=s, color=c),
        bottom = Side(style=s, color=c)
    )


def make_font(n='Segoe UI', s=10, b=False,
              i=False, c=BLACK, u=None):
    return Font(name=n, size=s, bold=b, italic=i, color=c, underline=u)


def make_align(h='center', v='center', w=False, i=0):
    return Alignment(horizontal=h, vertical=v, wrap_text=w, indent=i)


thin_border   = make_border(BORDER_BLUE, 'thin')
header_border = make_border(BORDER_DARK, 'thin')
medium_border = make_border(DARK_NAVY,   'medium')
fy_sep_border = Border(
    left   = Side(style='medium', color=FY_SEP_COLOR),
    right  = Side(style='thin',   color=BORDER_BLUE),
    top    = Side(style='thin',   color=BORDER_BLUE),
    bottom = Side(style='thin',   color=BORDER_BLUE)
)


def fill_row(ws, r, c, end_col):
    for col in range(1, end_col + 1):
        ws.cell(row=r, column=col).fill = make_fill(c)


# ============================================================
#                    FY HELPERS
# ============================================================

def get_fy_label(month_str):
    """Indian FY: Apr → Mar. e.g. Apr-2025 → FY25-26"""
    p = pd.Period(month_str, freq='M')
    y, mo = p.year, p.month
    if mo >= 4:
        return f"FY{str(y)[-2:]}-{str(y+1)[-2:]}"
    return f"FY{str(y-1)[-2:]}-{str(y)[-2:]}"


def build_fy_groups(unique_months):
    """Group months by FY, preserving chronological order."""
    fy_groups = {}
    for m in unique_months:
        fy = get_fy_label(m)
        fy_groups.setdefault(fy, []).append(m)
    return fy_groups


# ============================================================
#                    PIVOT BUILDER
# ============================================================

def build_pivots(crd_df, unique_months, fy_groups):
    """
    Build booking + revenue pivots and per-FY / overall summaries.
    Returns:
        booking_pivot, revenue_pivot,
        fy_summary  {fy: {'TotBk','TotRev','AvgBk','AvgRev','AvgTkt','NumMonths'}},
        overall_summary  {'TotBk','TotRev','AvgBk','AvgRev','AvgTkt'},
        entities_sorted
    """
    booking_pivot = crd_df.pivot_table(
        index   = ['Group Hub', 'Group Entity'],
        columns = 'Month',
        values  = 'BookingNo',
        aggfunc = 'nunique',
        fill_value = 0
    ).reindex(columns=unique_months, fill_value=0)

    revenue_pivot = (
        crd_df.pivot_table(
            index   = ['Group Hub', 'Group Entity'],
            columns = 'Month',
            values  = 'Revenue',
            aggfunc = 'sum',
            fill_value = 0
        ).reindex(columns=unique_months, fill_value=0)
        / 100_000.0
    )

    # Per-FY summaries
    fy_summary = {}
    for fy, months in fy_groups.items():
        n           = len(months)
        tot_bk      = booking_pivot[months].sum(axis=1)
        tot_rev     = revenue_pivot[months].sum(axis=1)
        tot_rev_r2  = tot_rev.round(2)
        fy_summary[fy] = {
            'TotBk'    : tot_bk,
            'TotRev'   : tot_rev_r2,
            'AvgBk'    : (tot_bk / n).round(1),
            'AvgRev'   : (tot_rev / n).round(2),
            # Use rounded Lakhs so data-row AvgTkt matches GT formula
            'AvgTkt'   : (tot_rev_r2 * 100_000).div(tot_bk)
                            .replace([np.inf, -np.inf], 0)
                            .fillna(0).round(0),
            'NumMonths': n
        }

    # Overall
    num_months_total  = len(unique_months)
    total_bk_all      = booking_pivot[unique_months].sum(axis=1)
    total_rev_all     = revenue_pivot[unique_months].sum(axis=1)
    total_rev_all_r2  = total_rev_all.round(2)
    overall_summary = {
        'TotBk' : total_bk_all,
        'TotRev': total_rev_all_r2,
        'AvgBk' : (total_bk_all / num_months_total).round(1),
        'AvgRev': (total_rev_all / num_months_total).round(2),
        'AvgTkt': (total_rev_all_r2 * 100_000).div(total_bk_all)
                     .replace([np.inf, -np.inf], 0)
                     .fillna(0).round(0)
    }

    # Sort entities by overall total revenue desc
    sort_idx      = total_rev_all.sort_values(ascending=False).index
    booking_pivot = booking_pivot.reindex(sort_idx)
    revenue_pivot = revenue_pivot.reindex(sort_idx)

    for fy in fy_summary:
        for k in ['TotBk', 'TotRev', 'AvgBk', 'AvgRev', 'AvgTkt']:
            fy_summary[fy][k] = fy_summary[fy][k].reindex(sort_idx)

    for k in ['TotBk', 'TotRev', 'AvgBk', 'AvgRev', 'AvgTkt']:
        overall_summary[k] = overall_summary[k].reindex(sort_idx)

    entities_sorted = list(sort_idx)
    return (booking_pivot, revenue_pivot,
            fy_summary, overall_summary, entities_sorted)


# ============================================================
#                    LAYOUT CALCULATOR
# ============================================================

def compute_layout(fy_groups, show_overall):
    """
    Compute Excel column positions for each FY and overall summary.
    Returns dict:
        fy_layout      {fy: {'month_start','month_end','sum_start','sum_end'}}
        overall_start, overall_end (or None)
        total_cols
    """
    FIXED_COLS   = 3
    SUMMARY_COLS = 5
    col_ptr      = FIXED_COLS + 1
    fy_layout    = {}

    for fy, months in fy_groups.items():
        n       = len(months)
        m_start = col_ptr
        m_end   = m_start + n * 2 - 1
        s_start = m_end + 1
        s_end   = s_start + SUMMARY_COLS - 1
        fy_layout[fy] = {
            'month_start': m_start, 'month_end': m_end,
            'sum_start'  : s_start, 'sum_end'  : s_end
        }
        col_ptr = s_end + 1

    if show_overall:
        overall_start = col_ptr
        overall_end   = overall_start + SUMMARY_COLS - 1
        total_cols    = overall_end
    else:
        overall_start = None
        overall_end   = None
        total_cols    = col_ptr - 1

    return fy_layout, overall_start, overall_end, total_cols


# ============================================================
#                    HEADER BUILDER
# ============================================================

def build_headers(ws, fy_groups, fy_summary, fy_layout,
                  overall_start, overall_end, show_overall,
                  total_cols, num_entities, num_hubs,
                  start_month, end_month, unique_months):
    """
    Row 1-2 : Title
    Row 3   : Subtitle
    Row 4   : Note
    Row 5   : Spacer
    Row 6   : FY group / Overall group headers
    Row 7   : Month sub-headers + summary sub-headers
    """

    fy_list = list(fy_groups.keys())

    # ── Title Rows (1-2) ─────────────────────────────────
    ws.merge_cells(
        start_row=1, start_column=1,
        end_row=2, end_column=total_cols
    )
    tc = ws.cell(
        1, 1,
        "CRD VERTICAL — MONTHLY BOOKING & REVENUE PERFORMANCE"
    )
    tc.font      = make_font(s=22, b=True, c=WHITE)
    tc.fill      = make_fill(DARK_NAVY)
    tc.alignment = make_align()
    fill_row(ws, 1, DARK_NAVY, total_cols)
    fill_row(ws, 2, DARK_NAVY, total_cols)
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 26

    # ── Subtitle Row (3) ─────────────────────────────────
    ws.merge_cells(
        start_row=3, start_column=1,
        end_row=3, end_column=total_cols
    )
    fy_str = " | ".join([
        f"{fy} ({fy_summary[fy]['NumMonths']}m)"
        for fy in fy_list
    ])
    sc_val = (
        f"Reporting Period: {start_month} to {end_month}  |  "
        f"Generated: {pd.Timestamp.now().strftime('%d-%b-%Y %H:%M')}  |  "
        f"Entities: {num_entities:,}  |  Hubs: {num_hubs}  |  {fy_str}"
    )
    sc = ws.cell(3, 1, sc_val)
    sc.font      = make_font(s=11, i=True, c=WHITE)
    sc.fill      = make_fill(SUBTITLE_BG)
    sc.alignment = make_align()
    ws.row_dimensions[3].height = 24

    # ── Note Row (4) ─────────────────────────────────────
    ws.merge_cells(
        start_row=4, start_column=1,
        end_row=4, end_column=total_cols
    )
    nc_val = (
        "⚑  Revenue in LAKHS INR  |  Per-FY summaries after each FY's months  |  "
        "Use filter on 'Group Hub' — Grand Total auto-updates"
    )
    nc = ws.cell(4, 1, nc_val)
    nc.font      = make_font(s=10.5, b=True, i=True, c=ACCENT_GOLD)
    nc.fill      = make_fill(NAVY_3)
    nc.alignment = make_align(w=True)
    ws.row_dimensions[4].height = 28

    # ── Spacer Row (5) ───────────────────────────────────
    fill_row(ws, 5, DARK_NAVY, total_cols)
    ws.row_dimensions[5].height = 3

    # ── Header Rows 6 (group) and 7 (sub) ────────────────
    month_grp_row = 6
    header_row    = 7
    ws.row_dimensions[month_grp_row].height = 28
    ws.row_dimensions[header_row].height    = 36

    # Fixed columns (merged across rows 6-7)
    fixed_headers = [
        ('S.No',         8),
        ('GROUP HUB',   11),
        ('GROUP ENTITY', 11)
    ]
    for c_idx, (lbl, sz) in enumerate(fixed_headers, 1):
        ws.merge_cells(
            start_row=month_grp_row, start_column=c_idx,
            end_row=header_row,      end_column=c_idx
        )
        c = ws.cell(month_grp_row, c_idx, lbl)
        c.font      = make_font(s=sz, b=True, c=WHITE)
        c.fill      = make_fill(HEADER_BG)
        c.alignment = make_align(w=True)
        c.border    = header_border

    # Per-FY headers
    for fy in fy_list:
        months  = fy_groups[fy]
        m_start = fy_layout[fy]['month_start']
        m_end   = fy_layout[fy]['month_end']
        s_start = fy_layout[fy]['sum_start']
        s_end   = fy_layout[fy]['sum_end']

        # Row 6: FY label over months
        ws.merge_cells(
            start_row=month_grp_row, start_column=m_start,
            end_row=month_grp_row,   end_column=m_end
        )
        fyc = ws.cell(month_grp_row, m_start, fy)
        fyc.font      = make_font(s=13, b=True, c=ACCENT_GOLD)
        fyc.fill      = make_fill(DARK_NAVY)
        fyc.alignment = make_align()
        fyc.border    = header_border
        for col in range(m_start, m_end + 1):
            ws.cell(month_grp_row, col).fill   = make_fill(DARK_NAVY)
            ws.cell(month_grp_row, col).border = header_border

        # Row 6: FY summary group label
        ws.merge_cells(
            start_row=month_grp_row, start_column=s_start,
            end_row=month_grp_row,   end_column=s_end
        )
        sfy = ws.cell(month_grp_row, s_start, f"{fy} SUMMARY")
        sfy.font      = make_font(s=12, b=True, c=ACCENT_AMBER)
        sfy.fill      = make_fill(SUMMARY_BG)
        sfy.alignment = make_align()
        sfy.border    = header_border
        for col in range(s_start, s_end + 1):
            ws.cell(month_grp_row, col).fill   = make_fill(SUMMARY_BG)
            ws.cell(month_grp_row, col).border = header_border

        # Row 7: Month sub-headers (Bookings / Revenue pairs)
        for m_idx, month in enumerate(months):
            cb = m_start + m_idx * 2
            cr = cb + 1
            b = ws.cell(header_row, cb, f"{month}\nBookings")
            b.font      = make_font(s=9, b=True, c=WHITE)
            b.fill      = make_fill(SUB_BOOK_BG)
            b.alignment = make_align(w=True)
            b.border    = header_border
            r = ws.cell(header_row, cr, f"{month}\nRevenue")
            r.font      = make_font(s=9, b=True, c=WHITE)
            r.fill      = make_fill(SUB_REV_BG)
            r.alignment = make_align(w=True)
            r.border    = header_border

        # Row 7: FY Summary sub-headers
        fy_sum_labels = [
            (s_start,     f"Total Bookings\n({fy})", True,  WHITE),
            (s_start + 1, f"Total Revenue\n({fy})",  True,  ACCENT_AMBER),
            (s_start + 2, f"Avg Bookings\n({fy})",   False, WHITE),
            (s_start + 3, f"Avg Revenue\n({fy})",    False, WHITE),
            (s_start + 4, f"Avg Ticket ₹\n({fy})",   False, WHITE),
        ]
        for col, lbl, bold, fc in fy_sum_labels:
            c = ws.cell(header_row, col, lbl)
            c.font      = make_font(s=9, b=bold, c=fc)
            c.fill      = make_fill(SUMMARY_BG)
            c.alignment = make_align(w=True)
            c.border    = header_border

    # Overall summary headers
    if show_overall:
        ws.merge_cells(
            start_row=month_grp_row, start_column=overall_start,
            end_row=month_grp_row,   end_column=overall_end
        )
        osc = ws.cell(month_grp_row, overall_start, "OVERALL SUMMARY")
        osc.font      = make_font(s=12, b=True, c=ACCENT_GOLD)
        osc.fill      = make_fill(OVERALL_BG)
        osc.alignment = make_align()
        osc.border    = header_border
        for col in range(overall_start, overall_end + 1):
            ws.cell(month_grp_row, col).fill   = make_fill(OVERALL_BG)
            ws.cell(month_grp_row, col).border = header_border

        overall_labels = [
            (overall_start,     "Total\nBookings",       True,  WHITE),
            (overall_start + 1, "Total\nRevenue",        True,  ACCENT_GOLD),
            (overall_start + 2, "Avg Bookings\n/ Month", False, WHITE),
            (overall_start + 3, "Avg Revenue\n/ Month",  False, WHITE),
            (overall_start + 4, "Avg Ticket\nSize (₹)",  False, WHITE),
        ]
        for col, lbl, bold, fc in overall_labels:
            c = ws.cell(header_row, col, lbl)
            c.font      = make_font(s=9.5, b=bold, c=fc)
            c.fill      = make_fill(OVERALL_BG)
            c.alignment = make_align(w=True)
            c.border    = header_border


# ============================================================
#                    DATA ROW WRITER
# ============================================================

def write_data_rows(ws, entities_sorted, booking_pivot, revenue_pivot,
                    fy_groups, fy_summary, overall_summary,
                    fy_layout, overall_start, show_overall,
                    data_start_row):
    """Write all entity data rows. Returns last data row."""

    fy_list = list(fy_groups.keys())

    for sno, (hub, entity) in enumerate(entities_sorted, 1):
        row     = data_start_row + sno - 1
        is_even = (sno % 2 == 0)

        row_bg     = ALT_ROW_2       if is_even else ALT_ROW_1
        bk_bg      = DATA_BOOK_BG2   if is_even else DATA_BOOK_BG1
        rev_bg     = DATA_REV_BG2    if is_even else DATA_REV_BG1
        fy_sum_bg  = FY_SUMMARY_BG2  if is_even else FY_SUMMARY_BG1
        overall_bg = OVERALL_COL_BG2 if is_even else OVERALL_COL_BG1

        ws.row_dimensions[row].height = 19

        # Fixed columns
        c = ws.cell(row, 1, sno)
        c.font      = make_font(s=9, c=GREY_MED)
        c.fill      = make_fill(row_bg)
        c.alignment = make_align()
        c.border    = thin_border

        c = ws.cell(row, 2, hub)
        c.font      = make_font(s=10, b=True, c=DARK_NAVY)
        c.fill      = make_fill(row_bg)
        c.alignment = make_align(h='left', i=1)
        c.border    = thin_border

        c = ws.cell(row, 3, entity)
        c.font      = make_font(s=10, c=BLACK)
        c.fill      = make_fill(row_bg)
        c.alignment = make_align(h='left', i=1)
        c.border    = thin_border

        # Per-FY month cells + FY summary cells
        for fy_idx, fy in enumerate(fy_list):
            months  = fy_groups[fy]
            m_start = fy_layout[fy]['month_start']
            s_start = fy_layout[fy]['sum_start']

            # Month cells
            for m_idx, month in enumerate(months):
                cb_col = m_start + m_idx * 2
                cr_col = cb_col + 1
                bval = booking_pivot.loc[(hub, entity), month]
                rval = revenue_pivot.loc[(hub, entity), month]

                b = ws.cell(row, cb_col)
                b.value         = int(bval)
                b.number_format = '#,##0;;"-"'
                b.font          = make_font(
                    s=10, b=(bval > 0),
                    c=DARK_NAVY if bval > 0 else GREY_MED
                )
                b.fill          = make_fill(bk_bg)
                b.alignment     = make_align()
                b.border        = (fy_sep_border
                                   if (m_idx == 0 and fy_idx > 0)
                                   else thin_border)

                r = ws.cell(row, cr_col)
                r.value         = float(rval)
                r.number_format = '#,##0.00;;"-"'
                r.font          = make_font(
                    s=10,
                    c=BLACK if rval > 0 else GREY_MED
                )
                r.fill          = make_fill(rev_bg)
                r.alignment     = make_align()
                r.border        = thin_border

            # FY summary cells
            fy_vals = [
                (fy_summary[fy]['TotBk'].loc[(hub, entity)],  '#,##0',    11, True,  GT_FONT_DARK),
                (fy_summary[fy]['TotRev'].loc[(hub, entity)], '#,##0.00', 11, True,  GT_FONT_DARK),
                (fy_summary[fy]['AvgBk'].loc[(hub, entity)],  '#,##0.0',  10, False, GREY_DARK),
                (fy_summary[fy]['AvgRev'].loc[(hub, entity)], '#,##0.00', 10, False, GREY_DARK),
                (fy_summary[fy]['AvgTkt'].loc[(hub, entity)], '#,##0',    10, False, GREY_DARK),
            ]
            for i, (val, fmt, sz, bold, fc) in enumerate(fy_vals):
                c = ws.cell(row, s_start + i)
                c.value         = float(val) if val is not None else 0
                c.number_format = fmt + ';;"-"'
                c.font          = make_font(s=sz, b=bold, c=fc)
                c.fill          = make_fill(fy_sum_bg)
                c.alignment     = make_align()
                c.border        = (fy_sep_border if i == 0 else thin_border)

        # Overall summary cells
        if show_overall:
            ov_vals = [
                (overall_summary['TotBk'].loc[(hub, entity)],  '#,##0',    12, True,  GT_FONT_DARK),
                (overall_summary['TotRev'].loc[(hub, entity)], '#,##0.00', 12, True,  GT_FONT_DARK),
                (overall_summary['AvgBk'].loc[(hub, entity)],  '#,##0.0',  10, False, GREY_DARK),
                (overall_summary['AvgRev'].loc[(hub, entity)], '#,##0.00', 10, False, GREY_DARK),
                (overall_summary['AvgTkt'].loc[(hub, entity)], '#,##0',    10, False, GREY_DARK),
            ]
            for i, (val, fmt, sz, bold, fc) in enumerate(ov_vals):
                c = ws.cell(row, overall_start + i)
                c.value         = float(val) if val is not None else 0
                c.number_format = fmt + ';;"-"'
                c.font          = make_font(s=sz, b=bold, c=fc)
                c.fill          = make_fill(overall_bg)
                c.alignment     = make_align()
                c.border        = (fy_sep_border if i == 0 else thin_border)

    data_end_row = data_start_row + len(entities_sorted) - 1
    return data_end_row


# ============================================================
#                    GRAND TOTAL ROW
# ============================================================

def write_grand_total_row(ws, gt_row, data_start_row, data_end_row,
                          fy_groups, fy_summary, fy_layout,
                          overall_start, show_overall,
                          num_months_total, total_cols):
    """
    SUBTOTAL(9,...) = SUM (respects filter)
    Avg Bookings = SUM(TotBk) / NumMonths
    Avg Revenue  = SUM(TotRev) / NumMonths
    Avg Ticket   = SUM(TotRev in Lakhs) * 100000 / SUM(TotBk)
    """
    ws.row_dimensions[gt_row].height = 32

    # Fixed labels
    fixed_labels = [
        ('GRAND',                 13, ACCENT_GOLD, 'center'),
        ('TOTAL',                 13, ACCENT_GOLD, 'left'),
        ('(Updates with Filter)',  9, GREY_MED,    'center')
    ]
    for i, (val, size, fc, al) in enumerate(fixed_labels, 1):
        c = ws.cell(gt_row, i, val)
        c.font      = make_font(s=size, b=True, i=(size == 9), c=fc)
        c.fill      = make_fill(DARK_NAVY)
        c.alignment = make_align(h=al, i=1 if al == 'left' else 0)
        c.border    = medium_border

    def write_subtotal(col, fmt, sz=11, bold=True, fc=WHITE):
        cL = get_column_letter(col)
        c  = ws.cell(
            gt_row, col,
            f"=SUBTOTAL(9,{cL}{data_start_row}:{cL}{data_end_row})"
        )
        c.number_format = fmt
        c.font          = make_font(s=sz, b=bold, c=fc)
        c.fill          = make_fill(DARK_NAVY)
        c.alignment     = make_align()
        c.border        = medium_border

    def write_gt_cell(col, formula, fmt, sz=10, bold=False, fc=WHITE):
        c = ws.cell(gt_row, col, formula)
        c.number_format = fmt
        c.font          = make_font(s=sz, b=bold, c=fc)
        c.fill          = make_fill(DARK_NAVY)
        c.alignment     = make_align()
        c.border        = medium_border

    # Per-FY: month subtotals + FY summary subtotals
    for fy, months in fy_groups.items():
        m_start = fy_layout[fy]['month_start']
        s_start = fy_layout[fy]['sum_start']
        n_months = fy_summary[fy]['NumMonths']

        # Month columns
        for m_idx in range(len(months)):
            write_subtotal(m_start + m_idx * 2,     '#,##0')
            write_subtotal(m_start + m_idx * 2 + 1, '#,##0.00')

        # FY Summary: Total Bookings + Total Revenue
        write_subtotal(s_start,     '#,##0',    sz=12, fc=ACCENT_GOLD)
        write_subtotal(s_start + 1, '#,##0.00', sz=12, fc=ACCENT_GOLD)

        # FY Summary: Avg Bookings = SUM(TotBk) / NumMonths
        cL_bk  = get_column_letter(s_start)
        cL_rev = get_column_letter(s_start + 1)
        write_gt_cell(
            s_start + 2,
            f"=IFERROR(SUBTOTAL(9,{cL_bk}{data_start_row}:{cL_bk}{data_end_row})"
            f"/{n_months},0)",
            '#,##0.0'
        )
        # FY Summary: Avg Revenue = SUM(TotRev) / NumMonths
        write_gt_cell(
            s_start + 3,
            f"=IFERROR(SUBTOTAL(9,{cL_rev}{data_start_row}:{cL_rev}{data_end_row})"
            f"/{n_months},0)",
            '#,##0.00'
        )
        # FY Summary: Avg Ticket = SUM(TotRev_Lakhs)*100000 / SUM(TotBk)
        write_gt_cell(
            s_start + 4,
            f"=IFERROR(SUBTOTAL(9,{cL_rev}{data_start_row}:{cL_rev}{data_end_row})"
            f"*100000/SUBTOTAL(9,{cL_bk}{data_start_row}:{cL_bk}{data_end_row}),0)",
            '#,##0'
        )

    # Overall summary subtotals
    if show_overall:
        write_subtotal(overall_start,     '#,##0',    sz=13, fc=ACCENT_GOLD)
        write_subtotal(overall_start + 1, '#,##0.00', sz=13, fc=ACCENT_GOLD)

        cL_bk  = get_column_letter(overall_start)
        cL_rev = get_column_letter(overall_start + 1)
        write_gt_cell(
            overall_start + 2,
            f"=IFERROR(SUBTOTAL(9,{cL_bk}{data_start_row}:{cL_bk}{data_end_row})"
            f"/{num_months_total},0)",
            '#,##0.0'
        )
        write_gt_cell(
            overall_start + 3,
            f"=IFERROR(SUBTOTAL(9,{cL_rev}{data_start_row}:{cL_rev}{data_end_row})"
            f"/{num_months_total},0)",
            '#,##0.00'
        )
        write_gt_cell(
            overall_start + 4,
            f"=IFERROR(SUBTOTAL(9,{cL_rev}{data_start_row}:{cL_rev}{data_end_row})"
            f"*100000/SUBTOTAL(9,{cL_bk}{data_start_row}:{cL_bk}{data_end_row}),0)",
            '#,##0'
        )

    # Fill remaining cells
    for col in range(1, total_cols + 1):
        cell = ws.cell(gt_row, col)
        if not cell.value:
            cell.fill   = make_fill(DARK_NAVY)
            cell.border = medium_border


# ============================================================
#                    COLUMN WIDTHS
# ============================================================

def set_column_widths(ws, fy_groups, fy_layout,
                      overall_start, show_overall):
    ws.column_dimensions['A'].width = 6
    ws.column_dimensions['B'].width = 22
    ws.column_dimensions['C'].width = 40

    for fy, months in fy_groups.items():
        m_start = fy_layout[fy]['month_start']
        s_start = fy_layout[fy]['sum_start']

        for m_idx in range(len(months)):
            ws.column_dimensions[
                get_column_letter(m_start + m_idx * 2)
            ].width = 11
            ws.column_dimensions[
                get_column_letter(m_start + m_idx * 2 + 1)
            ].width = 12

        for i, w in enumerate([13, 13, 12, 12, 13]):
            ws.column_dimensions[
                get_column_letter(s_start + i)
            ].width = w

    if show_overall:
        for i, w in enumerate([14, 14, 13, 13, 14]):
            ws.column_dimensions[
                get_column_letter(overall_start + i)
            ].width = w


# ============================================================
#                    PATH UTIL
# ============================================================

def resolve_path(relative_or_absolute_path):
    if not relative_or_absolute_path:
        return relative_or_absolute_path

    repo_root = get_repo_root()

    if os.path.isabs(relative_or_absolute_path):
        if os.path.exists(relative_or_absolute_path):
            return relative_or_absolute_path

        for marker in ['reports/', 'reports' + os.sep,
                       'data/', 'data' + os.sep]:
            idx = relative_or_absolute_path.find(marker)
            if idx != -1:
                rel_part = relative_or_absolute_path[idx:]
                return os.path.join(repo_root, rel_part)

    return os.path.join(repo_root, relative_or_absolute_path)


# ============================================================
#                       MAIN FUNCTION
# ============================================================

def main():
    print("\n" + "=" * 65)
    print("     📊 CRD BUSINESS REPORT GENERATOR (CI/CD)")
    print("=" * 65)

    # ── Step 1: Read State ────────────────────────────────
    print("\n[1/7] Reading state file...")
    state, state_path = read_state_file()

    combined_file_path = state['combined_file_path']
    output_dir         = state['output_dir']
    start_month        = state['start_month']
    end_month          = state['end_month']

    print(f"      Combined file : {combined_file_path}")
    print(f"      Output dir    : {output_dir}")
    print(f"      Date range    : {start_month} → {end_month}")

    # ── Step 2: Load & Filter ─────────────────────────────
    print("\n[2/7] Loading and filtering CRD data...")

    parquet_file_path = state.get('parquet_file_path', '')

    if parquet_file_path and \
       os.path.exists(parquet_file_path) and \
       parquet_file_path.endswith('.parquet'):
        print(f"      Reading parquet: {parquet_file_path}")
        df = pd.read_parquet(parquet_file_path, engine='pyarrow')
    elif combined_file_path and os.path.exists(combined_file_path):
        print(f"      Reading excel: {combined_file_path}")
        df = pd.read_excel(
            combined_file_path,
            sheet_name='Sheet1',
            engine='openpyxl'
        )
    else:
        raise FileNotFoundError(
            f"❌ No data file found.\n"
            f"   Parquet: {parquet_file_path}\n"
            f"   Excel  : {combined_file_path}"
        )

    crd_df = df[
        (df['Vertical'] == 'CRD') &
        (df['Group Entity'] != 'Not Found') &
        (df['Vertical']     != 'Not Found')
    ].copy()

    crd_df['Revenue']   = pd.to_numeric(
        crd_df['Revenue'], errors='coerce'
    ).fillna(0)
    crd_df['BookingNo'] = crd_df['BookingNo'].astype(str)

    for col in ['Group Hub', 'Group Entity', 'Vertical']:
        crd_df.loc[
            crd_df[col].isin(['', 'nan', 'None', 'NaN', 'NOT FOUND']),
            col
        ] = 'Unassigned'
    crd_df['Month'] = crd_df['Month'].astype(str).str.strip()

    # ── Step 3: Months + FY grouping ──────────────────────
    print("\n[3/7] Sorting months + detecting FYs...")

    def month_sort_key(m):
        try:
            return pd.Period(m, freq='M')
        except Exception:
            return pd.Period('1900-01', freq='M')

    unique_months = sorted(
        [m for m in crd_df['Month'].dropna().unique()
         if m not in ('', 'nan', '(blank)')],
        key=month_sort_key
    )
    num_months_total = len(unique_months)

    fy_groups    = build_fy_groups(unique_months)
    fy_list      = list(fy_groups.keys())
    show_overall = len(fy_list) > 1

    print(f"      Months : {unique_months[0]} → {unique_months[-1]} ({num_months_total})")
    print(f"      FYs    : {fy_list}")
    for fy in fy_list:
        print(f"        {fy}: {len(fy_groups[fy])} month(s) — "
              f"{fy_groups[fy][0]} to {fy_groups[fy][-1]}")

    # ── Step 4: Pivots + Summaries ────────────────────────
    print("\n[4/7] Building pivots + FY summaries...")
    (booking_pivot, revenue_pivot,
     fy_summary, overall_summary,
     entities_sorted) = build_pivots(
        crd_df, unique_months, fy_groups
    )

    num_entities = len(entities_sorted)
    num_hubs     = crd_df['Group Hub'].nunique()
    print(f"      Entities : {num_entities:,}")
    print(f"      Hubs     : {num_hubs}")

    # ── Step 5: Layout ────────────────────────────────────
    print("\n[5/7] Building Excel workbook...")
    (fy_layout, overall_start,
     overall_end, total_cols) = compute_layout(fy_groups, show_overall)

    data_start_row = 8

    # ── Step 6: Workbook ──────────────────────────────────
    wb = Workbook()
    ws = wb.active
    ws.title = "CRD Report"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation   = 'landscape'
    ws.page_setup.fitToWidth    = 1
    ws.page_setup.fitToHeight   = 0

    build_headers(
        ws            = ws,
        fy_groups     = fy_groups,
        fy_summary    = fy_summary,
        fy_layout     = fy_layout,
        overall_start = overall_start,
        overall_end   = overall_end,
        show_overall  = show_overall,
        total_cols    = total_cols,
        num_entities  = num_entities,
        num_hubs      = num_hubs,
        start_month   = start_month,
        end_month     = end_month,
        unique_months = unique_months
    )

    print(f"      Writing {num_entities:,} entity rows...")
    data_end_row = write_data_rows(
        ws              = ws,
        entities_sorted = entities_sorted,
        booking_pivot   = booking_pivot,
        revenue_pivot   = revenue_pivot,
        fy_groups       = fy_groups,
        fy_summary      = fy_summary,
        overall_summary = overall_summary,
        fy_layout       = fy_layout,
        overall_start   = overall_start,
        show_overall    = show_overall,
        data_start_row  = data_start_row
    )

    gt_row = data_end_row + 1
    write_grand_total_row(
        ws               = ws,
        gt_row           = gt_row,
        data_start_row   = data_start_row,
        data_end_row     = data_end_row,
        fy_groups        = fy_groups,
        fy_summary       = fy_summary,
        fy_layout        = fy_layout,
        overall_start    = overall_start,
        show_overall     = show_overall,
        num_months_total = num_months_total,
        total_cols       = total_cols
    )

    # AutoFilter — apply on sub-header row (row 7)
    ws.auto_filter.ref = (
        f"A7:{get_column_letter(total_cols)}{data_end_row}"
    )

    set_column_widths(ws, fy_groups, fy_layout,
                      overall_start, show_overall)

    # Freeze at first month column
    first_month_col = 4
    ws.freeze_panes = ws.cell(
        row=data_start_row, column=first_month_col
    )

    # ── Step 7: Save Report ───────────────────────────────
    print("\n[6/7] Saving report...")

    timestamp   = datetime.now().strftime("%Y%m%d_%H%M%S")
    report_file = os.path.join(
        output_dir,
        f"CRD_Business_Report_{timestamp}.xlsx"
    )

    has_not_found       = state.get("has_not_found", False)
    not_found_json_path = state.get("not_found_json_path", "")

    nf_df = pd.DataFrame()

    if has_not_found and not_found_json_path:
        if os.path.exists(not_found_json_path):
            try:
                with open(not_found_json_path, 'r',
                          encoding='utf-8') as f:
                    nf_records = json.load(f)

                nf_df = pd.DataFrame(nf_records)

                expected_cols = [
                    'Month', 'CorporateID', 'Corporate',
                    'Hub', 'Group Hub',
                    'Group Entity', 'Vertical',
                    'Bookings Count', 'Revenue'
                ]
                nf_df = nf_df[
                    [c for c in expected_cols if c in nf_df.columns]
                ]
                print(f"      ✅ Not Found data: {len(nf_df)} rows")
            except Exception as e:
                print(f"      ❌ Could not load JSON: {e}")
                nf_df = pd.DataFrame()
        else:
            print(f"      ⚠️  JSON not found at: {not_found_json_path}")

    wb.save(report_file)
    print(f"      ✅ Sheet 1 (CRD Report) saved.")

    # ── Sheet 2: Not Found ────────────────────────────────
    if has_not_found and not nf_df.empty:
        print(f"      📋 Writing Not Found Summary Sheet 2 "
              f"({len(nf_df)} rows)...")

        from openpyxl import load_workbook as lw
        from openpyxl.styles import (
            Font       as OFont,
            PatternFill as OFill,
            Alignment  as OAlign,
            Border     as OBorder,
            Side       as OSide
        )
        from openpyxl.utils import get_column_letter as gcl

        wb2   = lw(report_file)
        nf_ws = wb2.create_sheet(title="Not Found Summary")
        nf_ws.sheet_view.showGridLines = False

        C_TITLE  = "0D1B2A"
        C_HEADER = "2E5090"
        C_WHITE  = "FFFFFF"
        C_ALT1   = "F7FAFF"
        C_ALT2   = "E8F0FE"
        C_GOLD   = "FFF2CC"
        C_GOLDF  = "7F6000"
        C_BORDER = "B4C6E7"
        C_BLACK  = "1A1A1A"

        def xfill(c):
            return OFill(start_color=c, end_color=c, fill_type='solid')

        def xborder():
            s = OSide(style='thin', color=C_BORDER)
            return OBorder(left=s, right=s, top=s, bottom=s)

        def xfont(sz=10, bold=False, color=C_BLACK, italic=False):
            return OFont(
                name='Segoe UI', size=sz,
                bold=bold, color=color, italic=italic
            )

        def xalign(h='center', v='center', wrap=False):
            return OAlign(horizontal=h, vertical=v, wrap_text=wrap)

        nf_ws.merge_cells(start_row=1, start_column=1,
                          end_row=1, end_column=9)
        t = nf_ws.cell(
            1, 1,
            "NOT FOUND ENTITIES — GROUP ENTITY & VERTICAL MISSING"
        )
        t.font      = xfont(sz=14, bold=True, color=C_WHITE)
        t.fill      = xfill(C_TITLE)
        t.alignment = xalign()
        nf_ws.row_dimensions[1].height = 28

        nf_ws.merge_cells(start_row=2, start_column=1,
                          end_row=2, end_column=9)
        s = nf_ws.cell(
            2, 1,
            f"Total Records: {len(nf_df):,}  |  "
            f"Please fill Group Entity & Vertical columns "
            f"and upload to data/ folder."
        )
        s.font      = xfont(sz=10, italic=True, color=C_WHITE)
        s.fill      = xfill(C_HEADER)
        s.alignment = xalign()
        nf_ws.row_dimensions[2].height = 20

        headers = [
            'Month', 'CorporateID', 'Corporate',
            'Hub', 'Group Hub',
            'Group Entity', 'Vertical',
            'Bookings Count', 'Revenue'
        ]
        widths  = [12, 13, 40, 22, 15, 20, 15, 15, 14]

        for ci, (hdr, wd) in enumerate(zip(headers, widths), 1):
            c = nf_ws.cell(3, ci, hdr)
            c.font      = xfont(sz=11, bold=True, color=C_WHITE)
            c.fill      = xfill(C_HEADER)
            c.alignment = xalign(wrap=True)
            c.border    = xborder()
            nf_ws.column_dimensions[gcl(ci)].width = wd

        nf_ws.row_dimensions[3].height = 24

        HIGHLIGHT = [6, 7]
        for ri, (_, row) in enumerate(nf_df.iterrows(), 4):
            bg = C_ALT2 if (ri % 2 == 0) else C_ALT1
            for ci, col in enumerate(headers, 1):
                raw = row.get(col, '')
                if pd.isna(raw) or str(raw) == 'nan':
                    val = ''
                else:
                    val = raw
                if col in ['Bookings Count', 'Revenue']:
                    try:
                        val = int(float(str(val)))
                    except Exception:
                        val = 0

                cell = nf_ws.cell(ri, ci, val)
                if ci in HIGHLIGHT:
                    cell.fill = xfill(C_GOLD)
                    cell.font = xfont(sz=10, bold=True, color=C_GOLDF)
                else:
                    cell.fill = xfill(bg)
                    cell.font = xfont(sz=10)

                if ci in [1, 2, 5, 8, 9]:
                    cell.alignment = xalign(h='center')
                elif ci in [3, 4]:
                    cell.alignment = xalign(h='left')
                else:
                    cell.alignment = xalign(h='center')
                cell.border = xborder()
            nf_ws.row_dimensions[ri].height = 18

        nf_ws.freeze_panes    = nf_ws.cell(row=4, column=1)
        nf_ws.auto_filter.ref = f"A3:{gcl(9)}{3 + len(nf_df)}"

        for sname in wb2.sheetnames:
            if sname != "Not Found Summary":
                wb2[sname].title = "CRD Report"
                break

        wb2.save(report_file)
        print(f"      ✅ Final report saved with 2 sheets")
    else:
        print("      ✅ Report saved: CRD Report only")

    print(f"\n      📁 Saved at: {report_file}")

    # ── Update State ──────────────────────────────────────
    print("\n[7/7] Updating state file...")
    update_state_file(state_path, {
        "final_report_path"  : report_file,
        "total_entities"     : num_entities,
        "total_bookings"     : int(overall_summary['TotBk'].sum()),
        "total_revenue_lakhs": round(
            float(overall_summary['TotRev'].sum()), 2
        )
    })

    print("\n" + "=" * 65)
    print("     ✅ CRD BUSINESS REPORT — COMPLETE")
    print("=" * 65)


if __name__ == "__main__":
    main()
